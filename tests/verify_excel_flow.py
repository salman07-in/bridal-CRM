import sys
from pathlib import Path
from datetime import date
import io
import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from app import create_app
from app.config import Config
from app.extensions import db
from app.models.customer import Customer
from app.services.excel_service import ExcelService


def test_excel_roundtrip():
    print("[*] Testing Full Excel/CSV Import and Export Pipeline with Realistic Sample Data...")

    # 1. Create a mock Excel workbook in memory
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Register2021"

    # Header row
    ws.append(["Bride Name", "Mobile Number", "Wedding Date", "Makeup Date", "Package", "Amount", "Branch", "Consent"])
    # Sample records
    ws.append(["Rukhsar Bano", "+91 98765 11122", "14/02/2021", "14/02/2021", "Signature Bridal HD", 35000, "Main Studio", "Yes"])
    ws.append(["Heena Kausar", "09876533344", "20/11/2022", "19/11/2022", "Airbrush Royal Bridal", 40000, "City Centre Branch", "Y"])
    ws.append(["Nida Shaikh", "9876555566", "", "05/10/2023", "Classic Bridal Makeover", 25000, "Bridal Lounge", "No"]) # Missing wedding date
    ws.append(["Duplicate Test", "9876511122", "14/02/2021", "14/02/2021", "Signature Bridal HD", 35000, "Main Studio", "Yes"]) # Duplicate phone
    ws.append(["", "9876577788", "10/01/2023", "10/01/2023", "Package", 20000, "Main Studio", "Yes"]) # Missing name error

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    # 2. Parse workbook into rows
    rows = ExcelService.parse_file_to_rows(stream, "sample_register.xlsx")
    assert len(rows) == 5, f"Expected 5 data rows, got {len(rows)}"
    print(f"[+] Parsed {len(rows)} data rows from sample Excel file.")

    # 3. Validate & Preview
    app = create_app(Config)
    with app.app_context():
        report = ExcelService.validate_and_preview_import(rows)
        print(f"[+] Validation Preview: {report['valid_count']} Valid, {report['duplicate_count']} Duplicate(s), {report['invalid_count']} Error(s)")

        assert report["valid_count"] == 3
        assert report["duplicate_count"] == 1
        assert report["invalid_count"] == 1

        # 4. Import Valid Records
        imported = ExcelService.execute_import(report["valid_records"])
        print(f"[+] Successfully committed {imported} sample records to database.")

        # 5. Export to Excel (.xlsx)
        customers = Customer.query.filter(Customer.mobile_number.in_(["9876511122", "9876533344", "9876555566"])).all()
        excel_out = ExcelService.export_customers_to_excel(customers)
        assert excel_out is not None

        exported_wb = openpyxl.load_workbook(excel_out)
        exported_ws = exported_wb.active
        assert exported_ws.title == "Bridal Customers"
        assert exported_ws.max_row == 4 # Header + 3 records

        # 6. Export to CSV (.csv)
        csv_out = ExcelService.export_customers_to_csv(customers)
        assert "Rukhsar Bano" in csv_out
        assert "Heena Kausar" in csv_out
        assert "Nida Shaikh" in csv_out
        print("[+] Verified Excel and CSV export generation and data fidelity.")

        # Cleanup test records
        for c in customers:
            db.session.delete(c)
        db.session.commit()
        print("[+] Cleaned up test import records.")

    print("[+] SUCCESS: Excel import/export pipeline verified completely!")


if __name__ == "__main__":
    test_excel_roundtrip()
