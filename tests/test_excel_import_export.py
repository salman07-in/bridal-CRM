import io
from datetime import date
import openpyxl
from app.services.excel_service import ExcelService
from app.models.customer import Customer
from app.extensions import db


def test_flexible_date_parsing():
    assert ExcelService.parse_flexible_date("15/08/2021") == date(2021, 8, 15)
    assert ExcelService.parse_flexible_date("15-08-2021") == date(2021, 8, 15)
    assert ExcelService.parse_flexible_date("2021-08-15") == date(2021, 8, 15)
    assert ExcelService.parse_flexible_date("2021/08/15") == date(2021, 8, 15)
    assert ExcelService.parse_flexible_date(date(2021, 8, 15)) == date(2021, 8, 15)
    assert ExcelService.parse_flexible_date("") is None
    assert ExcelService.parse_flexible_date(None) is None
    assert ExcelService.parse_flexible_date("invalid-date-string") is None


def test_indian_mobile_normalization():
    # Valid formats
    phone, err = ExcelService.normalize_indian_mobile("9876543210")
    assert err is None and phone == "9876543210"

    phone, err = ExcelService.normalize_indian_mobile("+91 98765 43210")
    assert err is None and phone == "9876543210"

    phone, err = ExcelService.normalize_indian_mobile("09876543210")
    assert err is None and phone == "9876543210"

    # Invalid formats
    phone, err = ExcelService.normalize_indian_mobile("12345")
    assert phone is None and "must have 10 digits" in err

    phone, err = ExcelService.normalize_indian_mobile("5876543210") # Starts with 5
    assert phone is None and "start with 6, 7, 8, or 9" in err


def test_import_validation_and_duplicate_detection(app):
    with app.app_context():
        # Pre-seed existing customer
        existing = Customer(
            customer_code="SAF-EX-001",
            full_name="Existing Bride",
            mobile_number="9876500099",
        )
        db.session.add(existing)
        db.session.commit()

        # Batch with:
        # Row 1: Valid new record
        # Row 2: Duplicate with existing DB customer (9876500099)
        # Row 3: Duplicate inside same file
        # Row 4: Invalid (missing bride name)
        # Row 5: Invalid (bad phone number)
        rows = [
            {
                "Bride Name": "New Valid Bride",
                "Mobile": "9876500001",
                "Wedding Date": "22/10/2021",
                "Package": "Signature HD",
                "Amount": "25000",
                "Consent": "Yes",
            },
            {
                "Bride Name": "Conflict DB Bride",
                "Mobile": "9876500099", # In DB
                "Wedding Date": "10/05/2022",
            },
            {
                "Bride Name": "Duplicate In File",
                "Mobile": "9876500001", # Duplicate of Row 1 in file
            },
            {
                "Bride Name": "", # Missing name
                "Mobile": "9876500002",
            },
            {
                "Bride Name": "Bad Phone Bride",
                "Mobile": "1234", # Bad phone
            },
        ]

        report = ExcelService.validate_and_preview_import(rows)

        assert report["total_rows"] == 5
        assert report["valid_count"] == 1
        assert report["duplicate_count"] == 2 # 1 DB conflict + 1 in-file duplicate
        assert report["invalid_count"] == 2 # Missing name + bad phone

        assert report["valid_records"][0]["full_name"] == "New Valid Bride"
        assert report["valid_records"][0]["mobile_number"] == "9876500001"
        assert report["valid_records"][0]["wedding_date"] == date(2021, 10, 22)


def test_excel_and_csv_export(app):
    with app.app_context():
        c1 = Customer(
            customer_code="SAF-EXP-01",
            full_name="Export Bride 1",
            mobile_number="9876500011",
            wedding_date=date(2022, 4, 18),
            amount=32000.0,
            salon_branch="Main Studio",
        )
        c2 = Customer(
            customer_code="SAF-EXP-02",
            full_name="Export Bride 2",
            mobile_number="9876500012",
            wedding_date=None, # Missing date
            amount=15000.0,
        )
        db.session.add_all([c1, c2])
        db.session.commit()

        # Test Excel Export
        excel_stream = ExcelService.export_customers_to_excel([c1, c2])
        assert excel_stream is not None

        wb = openpyxl.load_workbook(excel_stream)
        ws = wb.active
        assert ws.title == "Bridal Customers"
        # Row 1 is header, Rows 2 and 3 are data
        assert ws.cell(row=2, column=1).value == "SAF-EXP-01"
        assert ws.cell(row=2, column=3).value == "Export Bride 1"
        assert ws.cell(row=3, column=6).value == "Missing (Review)"

        # Test CSV Export
        csv_text = ExcelService.export_customers_to_csv([c1, c2])
        assert "Export Bride 1" in csv_text
        assert "SAF-EXP-02" in csv_text


def test_import_and_export_financial_fields(app):
    """Test importing and exporting slip_number, category, advance paid, and balance due."""
    with app.app_context():
        rows = [
            {
                "Bride Name": "Financial Bride",
                "Mobile": "9876599988",
                "Slip Number": "SLIP-888",
                "Category": "Party Makeup",
                "Wedding Date": "10/12/2023",
                "Makeup Date": "10/12/2023",
                "Package": "Party Airbrush - Huda Beauty",
                "Amount": "5000",
                "Advance Paid": "2000",
                "Pre Makeup Reminder": "Yes",
                "Consent": "Yes",
            }
        ]

        report = ExcelService.validate_and_preview_import(rows)
        assert report["valid_count"] == 1
        record = report["valid_records"][0]
        assert record["slip_number"] == "SLIP-888"
        assert record["booking_type"] == "Party Makeup"
        assert record["amount"] == 5000.0
        assert record["advance_amount"] == 2000.0
        assert record["balance_due"] == 3000.0
        assert record["pre_makeup_reminder_enabled"] is True

        # Commit to DB
        imported = ExcelService.execute_import([record])
        assert imported == 1

        c = Customer.query.filter_by(mobile_number="9876599988").first()
        assert c is not None
        assert c.slip_number == "SLIP-888"
        assert c.booking_type == "Party Makeup"
        assert c.amount == 5000.0
        assert c.advance_amount == 2000.0
        assert c.balance_due == 3000.0
        assert c.payment_status == "Partial"

        # Export to CSV
        csv_text = ExcelService.export_customers_to_csv([c])
        assert "SLIP-888" in csv_text
        assert "Party Makeup" in csv_text
        assert "2000.00" in csv_text
        assert "3000.00" in csv_text
